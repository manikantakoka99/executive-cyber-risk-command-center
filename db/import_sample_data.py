#!/usr/bin/env python3
"""Idempotent importer: CyberCommandCenter_Sample_Data.xlsx → PostgreSQL.

Usage:
  .venv-ecc/bin/python db/import_sample_data.py [--xlsx PATH] [--database-url URL] [--reset]

--reset truncates all ECC tables before load (safe reload).
Compatible with ECC v2 foundation (RBAC + assets); maps workbook `users.role`
into user_roles and backfills assets from finding.affected_asset.
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl
import psycopg2
import psycopg2.extras

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = ROOT / "reference" / "CyberCommandCenter_Sample_Data.xlsx"
DEFAULT_URL = os.environ.get(
    "DATABASE_URL",
    "postgresql://ecc:ecc_dev_password@localhost:5433/cyber_command_center",
)

# Child → parent truncate order (v2 + MVP)
TABLE_ORDER = [
    "schema_migrations",  # keep — removed from truncate below
    "risk_verifications",
    "scenario_evidence",
    "scenario_controls",
    "risk_assessment_factors",
    "risk_assessments",
    "likelihood_assessments",
    "impact_assessments",
    "risk_scenario_signals",
    "risk_scenario_findings",
    "risk_scenario_assets",
    "risk_scenario_domains",
    "actions",
    "compliance_assessments",
    "framework_controls",
    "security_signals",
    "raw_events",
    "connector_sync_runs",
    "connectors",
    "asset_relationships",
    "asset_business_processes",
    "asset_identifiers",
    "assets",
    "business_processes",
    "business_units",
    "evidence",
    "controls",
    "risk_scenarios",
    "risk_tolerance_rules",
    "risk_policy_factors",
    "risk_policies",
    "role_permissions",
    "user_roles",
    "audit_log",
    "decision_action_log",
    "executive_decisions",
    "business_impact_assessments",
    "risk_findings",
    "domain_risk_snapshots",
    "enterprise_risk_score_snapshots",
    "org_compliance_status",
    "command_center_activity",
    "generated_reports",
    "org_domain_subscriptions",
    "users",
    "organizations",
    "risk_domains",
    "compliance_frameworks",
    # roles / permissions retained (system seed)
]

TRUNCATE_SKIP = {"schema_migrations", "roles", "permissions", "role_permissions"}

SHEET_LOAD_ORDER = [
    "organizations",
    "users",
    "risk_domains",
    "compliance_frameworks",
    "org_domain_subscriptions",
    "domain_risk_snapshots",
    "enterprise_risk_score_snapshots",
    "risk_findings",
    "business_impact_assessments",
    "executive_decisions",
    "decision_action_log",
    "org_compliance_status",
    "command_center_activity",
    "generated_reports",
]


def sheet_rows(wb: openpyxl.Workbook, name: str) -> tuple[list[str], list[dict[str, Any]]]:
    if name not in wb.sheetnames:
        raise SystemExit(f"Missing worksheet: {name}")
    ws = wb[name]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        return [], []
    if header is None or header[0] is None or str(header[0]).startswith("("):
        return [], []
    cols = [str(c) for c in header]
    out: list[dict[str, Any]] = []
    for r in rows_iter:
        if r is None or all(c is None for c in r):
            continue
        out.append(dict(zip(cols, r)))
    return cols, out


def coerce(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (datetime, date, Decimal, bool, int, float)):
        return value
    if isinstance(value, str):
        s = value.strip()
        if s == "":
            return None
        if s.startswith("{") or s.startswith("["):
            try:
                return psycopg2.extras.Json(json.loads(s))
            except json.JSONDecodeError:
                return s
        return s
    return value


def table_columns(cur, table: str) -> set[str]:
    cur.execute(
        """
        SELECT column_name FROM information_schema.columns
        WHERE table_schema='public' AND table_name=%s
        """,
        (table,),
    )
    return {r[0] for r in cur.fetchall()}


def reset_tables(cur) -> None:
    tables = [t for t in TABLE_ORDER if t not in TRUNCATE_SKIP]
    # Only truncate tables that exist
    cur.execute(
        """
        SELECT tablename FROM pg_tables WHERE schemaname='public'
        """
    )
    existing = {r[0] for r in cur.fetchall()}
    to_trunc = [t for t in tables if t in existing]
    if to_trunc:
        cur.execute(f"TRUNCATE {', '.join(to_trunc)} RESTART IDENTITY CASCADE;")


def upsert_rows(cur, table: str, cols: list[str], rows: list[dict[str, Any]]) -> int:
    if not rows:
        return 0
    existing_cols = table_columns(cur, table)
    cols = [c for c in cols if c in existing_cols]
    if not cols:
        return 0

    pk_map = {
        "organizations": ["org_id"],
        "users": ["user_id"],
        "risk_domains": ["risk_domain_id"],
        "compliance_frameworks": ["framework_id"],
        "org_domain_subscriptions": ["org_id", "risk_domain_id"],
        "domain_risk_snapshots": ["snapshot_id"],
        "enterprise_risk_score_snapshots": ["snapshot_id"],
        "risk_findings": ["finding_id"],
        "business_impact_assessments": ["impact_id"],
        "executive_decisions": ["decision_id"],
        "decision_action_log": ["action_log_id"],
        "org_compliance_status": ["status_id"],
        "command_center_activity": ["activity_id"],
        "generated_reports": ["report_id"],
        "audit_log": ["log_id"],
    }
    pks = pk_map[table]
    col_sql = ", ".join(cols)
    placeholders = ", ".join(["%s"] * len(cols))
    updates = ", ".join(f"{c}=EXCLUDED.{c}" for c in cols if c not in pks)
    if updates:
        sql = (
            f"INSERT INTO {table} ({col_sql}) VALUES ({placeholders}) "
            f"ON CONFLICT ({', '.join(pks)}) DO UPDATE SET {updates}"
        )
    else:
        sql = (
            f"INSERT INTO {table} ({col_sql}) VALUES ({placeholders}) "
            f"ON CONFLICT ({', '.join(pks)}) DO NOTHING"
        )

    count = 0
    for row in rows:
        values = [coerce(row.get(c)) for c in cols]
        if any(values[cols.index(pk)] is None for pk in pks if pk in cols):
            continue
        cur.execute(sql, values)
        count += 1
    return count


def assign_user_roles(cur, user_rows: list[dict[str, Any]]) -> int:
    n = 0
    for row in user_rows:
        uid = coerce(row.get("user_id"))
        oid = coerce(row.get("org_id"))
        role_code = coerce(row.get("role"))
        if not uid or not oid or not role_code:
            continue
        cur.execute(
            """
            INSERT INTO user_roles (user_id, role_id, org_id, is_primary)
            SELECT %s, r.role_id, %s, true
            FROM roles r WHERE r.code = %s
            ON CONFLICT (user_id, role_id) DO UPDATE SET is_primary = true, org_id = EXCLUDED.org_id
            """,
            (uid, oid, str(role_code)),
        )
        n += cur.rowcount
    return n


def backfill_assets_from_findings(cur) -> int:
    cur.execute(
        """
        INSERT INTO assets (org_id, name, asset_type, criticality, environment, internet_exposed, status)
        SELECT DISTINCT f.org_id, f.affected_asset,
          CASE
            WHEN lower(f.affected_asset) LIKE '%gateway%'
              OR lower(f.affected_asset) LIKE '%portal%'
              OR lower(f.affected_asset) LIKE '%api%'
              OR lower(f.affected_asset) LIKE '%checkout%'
              OR lower(f.affected_asset) LIKE '%service%' THEN 'application'::asset_type
            WHEN lower(f.affected_asset) LIKE '%infrastructure%' THEN 'host'::asset_type
            ELSE 'other'::asset_type
          END,
          CASE
            WHEN lower(f.affected_asset) LIKE '%payment%'
              OR lower(f.affected_asset) LIKE '%gateway%' THEN 'critical'::asset_criticality
            ELSE 'medium'::asset_criticality
          END,
          'production'::asset_environment,
          (lower(f.affected_asset) LIKE '%public%'
            OR lower(f.affected_asset) LIKE '%customer%'
            OR lower(f.affected_asset) LIKE '%online%'
            OR lower(f.affected_asset) LIKE '%gateway%'),
          'active'
        FROM risk_findings f
        WHERE f.affected_asset IS NOT NULL
          AND NOT EXISTS (
            SELECT 1 FROM assets a WHERE a.org_id = f.org_id AND a.name = f.affected_asset
          )
        """
    )
    inserted = cur.rowcount
    cur.execute(
        """
        UPDATE risk_findings f
        SET asset_id = a.asset_id,
            first_seen_at = COALESCE(f.first_seen_at, f.detected_at),
            last_seen_at = COALESCE(f.last_seen_at, COALESCE(f.resolved_at, f.updated_at, f.detected_at))
        FROM assets a
        WHERE f.affected_asset IS NOT NULL
          AND a.org_id = f.org_id
          AND a.name = f.affected_asset
        """
    )
    return inserted


def main() -> int:
    parser = argparse.ArgumentParser(description="Import ECC sample workbook into Postgres")
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--database-url", default=DEFAULT_URL)
    parser.add_argument("--reset", action="store_true", help="Truncate ECC tables before import")
    args = parser.parse_args()

    if not args.xlsx.exists():
        print(f"Workbook not found: {args.xlsx}", file=sys.stderr)
        return 1

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    conn = psycopg2.connect(args.database_url)
    conn.autocommit = False
    summary: dict[str, int] = {}

    try:
        with conn.cursor() as cur:
            if args.reset:
                print("Resetting tables…")
                reset_tables(cur)

            user_sheet_rows: list[dict[str, Any]] = []
            for sheet in SHEET_LOAD_ORDER:
                cols, rows = sheet_rows(wb, sheet)
                if not cols:
                    print(f"  {sheet}: skipped (empty)")
                    summary[sheet] = 0
                    continue
                if sheet == "users":
                    user_sheet_rows = rows
                n = upsert_rows(cur, sheet, cols, rows)
                summary[sheet] = n
                print(f"  {sheet}: upserted {n}")

            if user_sheet_rows:
                summary["user_roles"] = assign_user_roles(cur, user_sheet_rows)
                print(f"  user_roles: assigned {summary['user_roles']}")

            summary["assets_backfill"] = backfill_assets_from_findings(cur)
            print(f"  assets_backfill: {summary['assets_backfill']}")

            # Phase 2: sample workbook has 12 domains — restore controlled unmapped domain
            cur.execute(
                """
                INSERT INTO risk_domains (code, display_name, short_label, yvi_solution, description)
                VALUES (
                  'unmapped',
                  'Unmapped / Unknown Source',
                  'Unmapped',
                  NULL,
                  'Holding domain for signals that could not be mapped to a subscribed ECC domain'
                )
                ON CONFLICT (code) DO NOTHING
                """
            )
            summary["unmapped_domain"] = cur.rowcount
            print(f"  unmapped_domain: ensured")

            if "audit_log" in wb.sheetnames:
                cols, rows = sheet_rows(wb, "audit_log")
                if cols and rows:
                    summary["audit_log"] = upsert_rows(cur, "audit_log", cols, rows)
                    print(f"  audit_log: upserted {summary['audit_log']}")
                else:
                    summary["audit_log"] = 0
                    print("  audit_log: 0 rows")

        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()

    print("\nImport complete:")
    for k, v in summary.items():
        print(f"  {k}: {v}")
    print("\nNext: .venv-ecc/bin/python db/validate_import.py && .venv-ecc/bin/python db/validate_v2.py")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
