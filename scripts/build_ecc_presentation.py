#!/usr/bin/env python3
"""Generate ECC_R&D_and_Implementation_Presentation.pptx from verified project facts."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "ECC_R&D_and_Implementation_Presentation.pptx"
ASSETS = ROOT / "docs" / "presentation-assets"
REF_IMG = ASSETS / "reference-preview.jpg"
IMPL_IMG = ASSETS / "chrome-overview.png"
if not REF_IMG.exists():
    REF_IMG = ROOT / "reference" / "preview.jpg"

# Widescreen 16:9
W, H = Inches(13.333), Inches(7.5)

BG = RGBColor(0x0B, 0x0F, 0x14)
PANEL = RGBColor(0x12, 0x18, 0x20)
CYAN = RGBColor(0x37, 0xB6, 0xE8)
TEXT = RGBColor(0xE8, 0xED, 0xF2)
MUTED = RGBColor(0x94, 0xA3, 0xB0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GOOD = RGBColor(0x3F, 0xB9, 0x66)
HIGH = RGBColor(0xF0, 0x92, 0x3E)


def set_run(run, size=18, bold=False, color=TEXT, font="Calibri"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_bg(slide):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    shape.line.fill.background()
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG
    # send to back
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def add_accent_bar(slide):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), H)
    bar.line.fill.background()
    bar.fill.solid()
    bar.fill.fore_color.rgb = CYAN


def add_footer(slide, page: int, total: int):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(7.05), Inches(10), Inches(0.3))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = f"YVI DWAN · Executive Cyber Risk Command Center  ·  {page}/{total}"
    set_run(run, 11, color=MUTED)
    p.alignment = PP_ALIGN.LEFT


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_accent_bar(slide)
    # panel
    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.8), Inches(11.8), Inches(3.6)
    )
    panel.adjustments[0] = 0.05
    panel.line.fill.background()
    panel.fill.solid()
    panel.fill.fore_color.rgb = PANEL

    tf = slide.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(11), Inches(0.4)).text_frame
    r = tf.paragraphs[0].add_run()
    r.text = "YVI DWAN  ·  EXECUTIVE RISK INTELLIGENCE"
    set_run(r, 14, True, CYAN)

    tf = slide.shapes.add_textbox(Inches(1.1), Inches(2.6), Inches(11), Inches(1.2)).text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    r.text = "Executive Cyber Risk Command Center"
    set_run(r, 36, True, TEXT)

    tf = slide.shapes.add_textbox(Inches(1.1), Inches(3.9), Inches(11), Inches(1.0)).text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    r.text = (
        "R&D → Implementation → Production-readiness\n"
        "Risk → Impact → Decision → Action → Closed Loop"
    )
    set_run(r, 18, False, MUTED)

    tf = slide.shapes.add_textbox(Inches(1.1), Inches(5.7), Inches(11), Inches(0.4)).text_frame
    r = tf.paragraphs[0].add_run()
    r.text = "Data-backed MVP  ·  PostgreSQL · Fastify · React  ·  Sample dataset verified"
    set_run(r, 14, False, MUTED)


def section_title(slide, title: str, subtitle: str | None = None):
    tf = slide.shapes.add_textbox(Inches(0.55), Inches(0.35), Inches(12), Inches(0.55)).text_frame
    r = tf.paragraphs[0].add_run()
    r.text = title
    set_run(r, 28, True, TEXT)
    if subtitle:
        tf = slide.shapes.add_textbox(Inches(0.55), Inches(0.95), Inches(12), Inches(0.4)).text_frame
        tf.word_wrap = True
        r = tf.paragraphs[0].add_run()
        r.text = subtitle
        set_run(r, 14, False, MUTED)


def bullets(slide, items: list[str], left=0.55, top=1.5, width=12.0, height=5.2, size=17):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = f"•  {item}"
        set_run(run, size, False, TEXT)


def two_col_bullets(slide, left_items, right_items, left_title, right_title):
    # left panel
    for x, title, items in (
        (0.55, left_title, left_items),
        (6.95, right_title, right_items),
    ):
        hdr = slide.shapes.add_textbox(Inches(x), Inches(1.45), Inches(5.7), Inches(0.35)).text_frame
        r = hdr.paragraphs[0].add_run()
        r.text = title
        set_run(r, 16, True, CYAN)
        bullets(slide, items, left=x, top=1.9, width=5.7, height=4.6, size=15)


def content_slide(prs, title, subtitle, items, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_accent_bar(slide)
    section_title(slide, title, subtitle)
    bullets(slide, items)
    add_footer(slide, page, total)
    return slide


def image_slide(prs, title, subtitle, image_path: Path, page, total, caption: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_accent_bar(slide)
    section_title(slide, title, subtitle)
    img_left, img_top = Inches(0.7), Inches(1.4)
    max_w, max_h = Inches(12.0), Inches(5.0)
    if image_path.exists():
        from PIL import Image

        with Image.open(image_path) as im:
            iw, ih = im.size
        aspect = iw / ih
        # fit inside max box
        box_aspect = max_w / max_h
        if aspect > box_aspect:
            w = max_w
            h = int(max_w / aspect)
        else:
            h = max_h
            w = int(max_h * aspect)
        # center horizontally in content area
        left = img_left + (max_w - w) // 2
        slide.shapes.add_picture(str(image_path), left, img_top, width=w, height=h)
    else:
        bullets(slide, [f"Image missing: {image_path.name}"], top=2.0)
    cap = slide.shapes.add_textbox(Inches(0.7), Inches(6.55), Inches(12), Inches(0.35)).text_frame
    r = cap.paragraphs[0].add_run()
    r.text = caption
    set_run(r, 12, False, MUTED)
    add_footer(slide, page, total)
    return slide


def table_slide(prs, title, subtitle, headers, rows, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_accent_bar(slide)
    section_title(slide, title, subtitle)
    cols = len(headers)
    table_shape = slide.shapes.add_table(
        len(rows) + 1, cols, Inches(0.55), Inches(1.55), Inches(12.2), Inches(4.8)
    )
    table = table_shape.table
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            for run in p.runs:
                set_run(run, 13, True, WHITE)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0x1B, 0x4D, 0x66)
    for r_i, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.cell(r_i, c)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    set_run(run, 12, False, TEXT)
            cell.fill.solid()
            cell.fill.fore_color.rgb = PANEL if r_i % 2 else RGBColor(0x16, 0x1E, 0x27)
    add_footer(slide, page, total)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    total = 20

    title_slide(prs)

    # 2 agenda
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_accent_bar(slide)
    section_title(slide, "Agenda", "From problem framing to a demonstrable MVP")
    two_col_bullets(
        slide,
        [
            "Problem & product principle",
            "Source materials consumed",
            "Reference vs implemented UI",
            "Architecture & data model",
            "Sample data & KPI derivation",
        ],
        [
            "Executive journeys & decisions",
            "Security / tenancy",
            "What shipped + test evidence",
            "Limitations & next phase",
            "How to run locally",
        ],
        "Discovery & design",
        "Build & readiness",
    )
    add_footer(slide, 2, total)

    content_slide(
        prs,
        "The problem",
        "CXOs have many technical dashboards — not one business-context command view",
        [
            "Domain tools (SOC, IAM, CSPM, VAPT, compliance…) produce signals in isolation",
            "Board packs are assembled manually, late, and in technical language",
            "Executives need residual risk, financial exposure, compliance, and decisions in one place",
            "Success metric: time-to-executive-insight measured in seconds, not hours",
        ],
        3,
        total,
    )

    content_slide(
        prs,
        "Product principle",
        "ECC is an aggregation layer above 12 YVI domain solutions — not a replacement",
        [
            "Risk → Impact → Decision → Action → Closed Loop",
            "Ingest / normalize / correlate domain signals into executive language",
            "Translate findings into financial & compliance impact",
            "Escalate only what exceeds executive tolerance into a decision queue",
            "Record actions immutably and feed closed-loop reporting",
        ],
        4,
        total,
    )

    content_slide(
        prs,
        "Source materials (authoritative)",
        "Implementation followed supplied contracts — not invented models",
        [
            "HTML mockup + preview.jpg — visual / information-architecture baseline",
            "High-Level Flow PDF — end-to-end product stages",
            "Database ERD PDF + cyber_command_center_schema.sql — domain model & RLS intent",
            "CyberCommandCenter_Sample_Data.xlsx — synthetic multi-tenant dataset to import",
            "Prior repo demo (5-domain hardcoded SPA) treated as superseded prototype",
        ],
        5,
        total,
    )

    image_slide(
        prs,
        "Reference UI baseline",
        "Static HTML / preview — visual target, not hardcoded production values",
        REF_IMG,
        6,
        total,
        "Note: mockup KPIs (e.g. score 62) differ from sample DB (Al Dhabi ≈ 42.83).",
    )

    image_slide(
        prs,
        "Implemented Command Center",
        "Data-backed Overview — KPIs computed from PostgreSQL for the selected organization",
        IMPL_IMG,
        7,
        total,
        "Decisions Awaiting Executive Action is the primary CXO interaction (inline Approve / Request Info).",
    )

    content_slide(
        prs,
        "Architecture (smallest viable)",
        "Reuse existing React/Vite app; add Postgres + Fastify aggregation API",
        [
            "PostgreSQL 16 via Docker Compose (host port 5433)",
            "Idempotent Excel → DB importer + FK validation report",
            "Fastify /api/ecc/* with org-scoped aggregation (overview endpoint)",
            "React UI: Overview, Domains, Findings, Decisions, Compliance, Reports, Audit",
            "Demo auth headers X-Org-Id / X-User-Id (SSO deferred)",
            "Explicitly deferred: microservices, K8s, WebSockets, FAIR, live external feeds",
        ],
        8,
        total,
    )

    content_slide(
        prs,
        "Data model (schema-faithful)",
        "15 tables · multi-tenant org_id · Risk→Impact→Decision linkage",
        [
            "Tenancy: organizations, users (cxo / board / program_office / analyst / auditor)",
            "Domains: risk_domains (12), org_domain_subscriptions",
            "Scores: domain_risk_snapshots, enterprise_risk_score_snapshots",
            "Evidence: risk_findings → business_impact_assessments → executive_decisions",
            "Closed loop: decision_action_log, command_center_activity, generated_reports, audit_log",
            "Compliance: compliance_frameworks + org_compliance_status",
        ],
        9,
        total,
    )

    table_slide(
        prs,
        "Sample dataset (imported)",
        "Workbook mapped 1:1 to schema — validation PASSED after reload",
        ["Entity", "Rows", "Notes"],
        [
            ["organizations", "5", "Incl. Al Dhabi, Falcon, Burj, Meridian, Zayed"],
            ["risk_domains", "12", "Full YVI portfolio taxonomy"],
            ["domain_risk_snapshots", "360", "6 months × 12 × 5"],
            ["enterprise snapshots", "30", "Trend source"],
            ["findings / impacts", "15 / 15", "Linked AED exposure"],
            ["executive_decisions", "15", "Awaiting varies by org"],
            ["compliance status", "50", "5 frameworks × orgs/snapshots"],
            ["activity / reports", "34 / 10", "Closed-loop artifacts"],
        ],
        10,
        total,
    )

    table_slide(
        prs,
        "KPI calculation model",
        "Server-side only — no hardcoded HTML prototype numbers in the UI",
        ["KPI", "Source", "Rule"],
        [
            ["Enterprise risk score", "enterprise_risk_score_snapshots", "Latest snapshot + MoM delta"],
            ["Business impact exposure", "business_impact_assessments", "Sum AED for open findings"],
            ["Compliance posture", "org_compliance_status", "Avg latest coverage %"],
            ["Decisions awaiting CXO", "executive_decisions", "status = awaiting_decision ONLY"],
            ["Risk by domain", "domain_risk_snapshots", "Latest per subscribed domain"],
            ["Risk trend", "enterprise snapshots", "Chronological score series"],
        ],
        11,
        total,
    )

    content_slide(
        prs,
        "Executive journeys (MVP)",
        "Designed for <30s posture comprehension and actionable follow-through",
        [
            "J1 CXO opens Overview → score, trend, exposure, compliance, awaiting queue",
            "J2 Domain card → score history, findings, linked decisions, exposure rollup",
            "J3 Finding → financial / downtime / BU / compliance scope → linked decision",
            "J4 Approve or Request More Information inline → action log + activity + audit",
            "Org switcher: Al Dhabi (0 awaiting, valid empty) · Falcon (2 awaiting, demo path)",
        ],
        12,
        total,
    )

    content_slide(
        prs,
        "Decision workflow",
        "Immutable history · explicit statuses · overview matches KPI",
        [
            "States: awaiting_decision · info_requested · approved · declined · closed",
            "Overview KPI & queue count only awaiting_decision (info_requested excluded)",
            "Primary overview actions: Approve (from recommended_action) + Request More Information",
            "Every mutation writes decision_action_log + audit_log + command_center_activity",
            "RBAC: CXO approve/decline · CXO/Program Office request-info/close · others read",
            "Duplicate submission guarded in UI (busy lock) + terminal-state 409 on API",
        ],
        13,
        total,
    )

    content_slide(
        prs,
        "Security & tenancy",
        "Frontend hiding is not isolation — API enforces organization context",
        [
            "All tenant routes require validated X-Org-Id + X-User-Id membership",
            "Findings, decisions, compliance, reports, audit filtered by org_id",
            "Cross-org decision read/mutate returns 404; analyst approve returns 403",
            "Domain UUIDs are shared catalog; findings/snapshots remain org-scoped (tested)",
            "Schema includes RLS policies; MVP app role still bypasses RLS — next-phase hardening",
            "Demo auth is intentional for workshops — not a production IdP",
        ],
        14,
        total,
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_accent_bar(slide)
    section_title(slide, "What we shipped", "MUST-HAVE flow complete · production-readiness pass applied")
    two_col_bullets(
        slide,
        [
            "DB + sample import/validation",
            "Aggregation APIs + overview",
            "Domains / findings / decisions",
            "Inline CXO decision actions",
            "Compliance · reports · audit",
            "Org switch + role-visible UI",
        ],
        [
            "RIDA chain on overview cards",
            "Honest polled freshness label",
            "Domain drill-down enrichment",
            "Empty/error/loading states",
            "Tenant isolation tests",
            "Typecheck + production build",
        ],
        "Core product",
        "Hardening",
    )
    add_footer(slide, 15, total)

    content_slide(
        prs,
        "Test evidence",
        "Gates executed during R&D and hardening",
        [
            "Sample import validation — row counts + FK checks PASSED",
            "Overview golden tests (Al Dhabi KPIs · Falcon queue ≡ KPI) PASSED",
            "Security tests (finding isolation · report scoping · approve/audit) PASSED",
            "E2E Risk→Impact→Decision→Action script PASSED",
            "Frontend TypeScript build (tsc + vite build) PASSED",
        ],
        16,
        total,
    )

    content_slide(
        prs,
        "Known limitations",
        "Explicit so stakeholders do not confuse MVP with production enterprise ops",
        [
            "Header demo auth — not SSO / MFA",
            "DB owner connection bypasses Postgres RLS (API filters active today)",
            "No live domain feed adapters yet (workbook is the source)",
            "Sample compliance_scope_impact often null — UI shows when present",
            "Board briefing is Markdown from live aggregates (PDF later)",
            "Al Dhabi sample has zero awaiting decisions — use Falcon for action demo",
        ],
        17,
        total,
    )

    content_slide(
        prs,
        "Recommended next phase",
        "Stay on current architecture — deepen trust, feeds, and board packaging",
        [
            "Phase A: SSO + connect as ccc_app with app.current_org_id RLS enforced",
            "Phase B: Wire 1–3 live domain feeds into findings/snapshots",
            "Phase C: Per-org risk appetite config (replace env default 50)",
            "Phase D: Decision aging / SLA + PDF board briefing",
            "Phase E: Playwright UI smoke (org switch + approve path)",
            "Do not add FAIR Monte Carlo, WebSockets, or microservices unless required",
        ],
        18,
        total,
    )

    content_slide(
        prs,
        "How to run locally",
        "Workshop / demo path",
        [
            "docker compose up -d   # Postgres on localhost:5433",
            ".venv-ecc/bin/python db/import_sample_data.py --reset && validate_import.py",
            "cd server && npm run dev   # API :4000",
            "cd web && npm run dev      # UI :5173 (proxies /api)",
            "Switch org to Falcon National Bank to exercise awaiting CXO decisions",
            "Docs: docs/ecc-*.md · checklist · production-readiness",
        ],
        19,
        total,
    )

    # close
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_accent_bar(slide)
    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(2.2), Inches(10.9), Inches(3.0)
    )
    panel.adjustments[0] = 0.05
    panel.line.fill.background()
    panel.fill.solid()
    panel.fill.fore_color.rgb = PANEL
    tf = slide.shapes.add_textbox(Inches(1.6), Inches(2.6), Inches(10), Inches(2.0)).text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    r.text = "One command view for cyber risk —\nin business language, with decisions that close the loop."
    set_run(r, 26, True, TEXT)
    tf2 = slide.shapes.add_textbox(Inches(1.6), Inches(4.5), Inches(10), Inches(0.5)).text_frame
    r = tf2.paragraphs[0].add_run()
    r.text = "Questions · pilot scoping · feed mapping workshop"
    set_run(r, 16, False, CYAN)
    add_footer(slide, 20, total)

    prs.save(OUT)
    return OUT


def inspect_overflow(path: Path):
    """Brief text overflow / density check for QA reporting."""
    prs = Presentation(str(path))
    issues = []
    for i, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for p in shape.text_frame.paragraphs:
                text = "".join(run.text for run in p.runs)
                if len(text) > 220:
                    issues.append(f"Slide {i}: long paragraph ({len(text)} chars)")
                for run in p.runs:
                    if run.font.size and run.font.size < Pt(10):
                        issues.append(f"Slide {i}: font <10pt")
    # picture presence
    has_pic = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type is not None and "PICTURE" in str(shape.shape_type):
                has_pic += 1
    return {
        "slides": len(prs.slides),
        "pictures": has_pic,
        "issues": issues[:12],
        "path": str(path),
    }


if __name__ == "__main__":
    out = build()
    report = inspect_overflow(out)
    print("WROTE", out)
    print("REPORT", report)
